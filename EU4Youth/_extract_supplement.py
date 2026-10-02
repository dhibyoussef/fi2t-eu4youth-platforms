# -*- coding: utf-8 -*-
import json
import re
from pathlib import Path
from docx import Document

OUT = Path(__file__).with_name("_extract_supplement.json")

BASE = Path(r"c:\Users\youss\OneDrive\Attachments\Desktop\fi2t-live-editor\EU4Youth\EU4Youth")
SITE = BASE / "site web EU4Youth.org-20260811T205332Z-1-001" / "site web EU4Youth.org"


def docx_text(path):
    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def main():
    result = {}
    p = SITE / "Cadrage_UXUI_Actualites_Opportunites_Glossaire_Publications (2).docx"
    result["cadrage_uxui"] = docx_text(p)

    html = (SITE / "Glossaire" / "glossaire_eu4youth_complet.html").read_text(
        encoding="utf-8", errors="ignore"
    )
    terms = re.findall(r'class="entry-term"[^>]*>([^<]+)<', html)
    sections = re.findall(
        r'section-header[^>]*>.*?<h2[^>]*>([^<]+)</h2>', html, re.S
    )
    tabs = re.findall(r'class="tab-btn[^"]*"[^>]*>([^<]+)<', html)
    result["glossaire"] = {
        "term_count": len(terms),
        "unique_terms": len(set(t.strip() for t in terms)),
        "sections": sections,
        "tabs": tabs,
        "sample_terms": [t.strip() for t in terms[:20]],
    }

    p2 = SITE / "EU4Youth org - Cahier des charges 2026 (2).docx"
    t2 = docx_text(p2)
    cahier_sections = {}
    for marker in [
        "2.1 Arborescence",
        "4.3 Filtres",
        "3.6 Opportunités",
        "3.8 Publications",
        "3.10 Actualités",
        "3.12 Glossaire",
        "6.1 Fonctionnalités",
        "Annexe A",
    ]:
        idx = t2.find(marker)
        if idx >= 0:
            cahier_sections[marker] = t2[idx : idx + 3000]
    result["cahier_sections"] = cahier_sections

    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
