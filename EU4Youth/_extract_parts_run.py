# -*- coding: utf-8 -*-
"""Re-extract correct Cadrage + dump key texts for synthesis."""
from __future__ import annotations
import json
import re
from pathlib import Path
from docx import Document
import pymupdf

BASE = Path(r"c:\Users\youss\OneDrive\Attachments\Desktop\fi2t-live-editor\EU4Youth\EU4Youth")
SITE = BASE / "site web EU4Youth.org-20260811T205332Z-1-001" / "site web EU4Youth.org"
OUT_DIR = Path(r"c:\Users\youss\OneDrive\Attachments\Desktop\fi2t-live-editor\EU4Youth\_extract_parts")
OUT_DIR.mkdir(exist_ok=True)


def docx_full(path: Path) -> str:
    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        style = (para.style.name if para.style else "") or ""
        t = para.text.strip()
        if not t:
            continue
        if "Heading" in style or "Title" in style:
            parts.append(f"\n## [{style}] {t}\n")
        else:
            parts.append(t)
    for ti, table in enumerate(doc.tables):
        parts.append(f"\n### TABLE {ti+1}\n")
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def pdf_all_text(path: Path, max_pages: int = 8) -> str:
    doc = pymupdf.open(str(path))
    chunks = [f"PAGE_COUNT={doc.page_count}"]
    for i in range(min(max_pages, doc.page_count)):
        chunks.append(f"\n===== PAGE {i+1} =====\n")
        chunks.append(doc[i].get_text("text"))
    doc.close()
    return "\n".join(chunks)


targets = {
    "01_cahier": SITE / "EU4Youth org - Cahier des charges 2026 (2).docx",
    "02_cadrage_uxui": SITE / "Cadrage_UXUI_Actualites_Opportunites_Glossaire_Publications (2).docx",
    "03_homepage": SITE / "Accueil" / "EU4Youth - Homepage fr.docx",
    "04_a_propos": list((SITE / "Le programme EU4Youth").glob("*.docx"))[0],
    "05_contact": list((SITE / "Formulaire de contact").glob("*.docx"))[0],
    "06_programme": list(BASE.glob("PROGRAMME*.docx"))[0],
    "07_sans_titre": BASE / "Document sans titre.docx",
    "08_ue_tunisie": list(BASE.glob("*Union*.docx"))[0],
}

meta = {}
for key, path in targets.items():
    text = docx_full(path)
    out = OUT_DIR / f"{key}.txt"
    out.write_text(text, encoding="utf-8")
    meta[key] = {"file": path.name, "chars": len(text), "out": str(out)}
    print(key, path.name, len(text))

# glossaire
html = (SITE / "Glossaire" / "glossaire_eu4youth_complet.html").read_text(encoding="utf-8", errors="ignore")
# count terms: look for pattern Terme / Définition structure
# From typical complete glossary HTML
terms_dt = re.findall(r"<dt[^>]*>(.*?)</dt>", html, re.I | re.S)
terms_h = re.findall(r"<h[2-4][^>]*>(.*?)</h[2-4]>", html, re.I | re.S)
# also data-term or .glossary-term
terms_div = re.findall(r'class=["\'][^"\']*(?:term|entry|mot)[^"\']*["\'][^>]*>(.*?)<', html, re.I | re.S)
# letter sections + definitions
# Try structured: <b>TERM</b> or <strong>
strongs = [re.sub(r"<[^>]+>", "", s).strip() for s in re.findall(r"<strong[^>]*>(.*?)</strong>", html, re.I | re.S)]
# Many glossaries use <p><b>Term</b> : def
bold_p = re.findall(r"<p[^>]*>\s*<(?:b|strong)[^>]*>(.*?)</(?:b|strong)>", html, re.I | re.S)
clean = []
for s in bold_p or strongs or terms_dt or terms_h:
    t = re.sub(r"<[^>]+>", "", s).strip()
    t = re.sub(r"\s+", " ", t)
    if 1 < len(t) < 100 and t.lower() not in {"glossaire", "eu4youth", "définition", "definition"}:
        clean.append(t)
seen = set()
uniq = []
for t in clean:
    k = t.lower()
    if k not in seen:
        seen.add(k)
        uniq.append(t)
# fallback: count occurrences of pattern "— " or " : " after bold
(OUT_DIR / "09_glossaire_meta.json").write_text(
    json.dumps(
        {
            "html_chars": len(html),
            "dt": len(terms_dt),
            "h": len(terms_h),
            "bold_p": len(bold_p),
            "strongs": len(strongs),
            "term_estimate": len(uniq),
            "sample": uniq[:40],
            "snippet_start": html[:1500],
        },
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)
print("glossaire terms", len(uniq))

# UI PDFs first pages
UI = BASE / "UI Web Design-20260807T094637Z-1-001" / "UI Web Design"
ui_meta = {}
for p in sorted(UI.glob("*.pdf")):
    text = pdf_all_text(p, max_pages=3)
    safe = re.sub(r"[^\w\-]+", "_", p.stem)[:60]
    (OUT_DIR / f"ui_{safe}.txt").write_text(text, encoding="utf-8")
    ui_meta[p.name] = {"pages": int(re.search(r"PAGE_COUNT=(\d+)", text).group(1)), "chars": len(text)}
    print("UI", p.name, ui_meta[p.name]["pages"])

# root projet pdfs
for p in BASE.glob("projet*.pdf"):
    text = pdf_all_text(p, max_pages=3)
    (OUT_DIR / f"ui_root_{p.stem}.txt").write_text(text, encoding="utf-8")
    print("ROOT", p.name)

# projects one-liners
proj_summaries = []
for folder in sorted((SITE / "Projets").iterdir()):
    if not folder.is_dir():
        continue
    docs = list(folder.glob("*.docx"))
    if not docs:
        continue
    text = docx_full(docs[0])
    (OUT_DIR / f"proj_{folder.name}.txt").write_text(text[:8000], encoding="utf-8")
    lines = [ln.strip() for ln in text.splitlines() if ln.strip() and not ln.startswith("#")]
    first = lines[0] if lines else folder.name
    para = next((ln for ln in lines if len(ln) > 50), "")[:300]
    proj_summaries.append({"folder": folder.name, "doc": docs[0].name, "first": first, "para": para, "lines": lines[:25]})

(OUT_DIR / "00_meta.json").write_text(
    json.dumps({"targets": meta, "ui": ui_meta, "projects": proj_summaries}, ensure_ascii=False, indent=2),
    encoding="utf-8",
)
print("DONE", OUT_DIR)
