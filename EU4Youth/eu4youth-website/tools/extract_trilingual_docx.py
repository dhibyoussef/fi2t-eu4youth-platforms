#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Parse trilingual EU4Youth Word docs into structured JSON under src/i18n/extracted/."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT.parent / "EU4Youth" / "Sections trilingues"
OUT_DIR = ROOT / "src" / "i18n" / "extracted"

ARABIC_RE = re.compile(r"[\u0600-\u06FF]")
KPI_RE = re.compile(
    r"^(\+?\d[\d\s,\.]*(?:\s*[%€])?|\d+\s*[%])\s*(.*)$", re.UNICODE
)
URL_RE = re.compile(r"https?://[^\s\]]+")

SECTION_SKIP = re.compile(
    r"^(Traitement graphique|Graphic treatment|المعالجة البصرية)",
    re.I,
)

HEADING_OBJECTIVE = re.compile(
    r"(objectif\s+g[ée]n[ée]ral|overall\s+objective|main\s+objective|"
    r"objectif\s+principal|الهدف\s+ال(?:رئيسي|عام))",
    re.I,
)
HEADING_SPECIFIC = re.compile(
    r"(objectifs?\s+sp[ée]cifiques?|specific\s+objectives?|الأهد(?:اف|ف)\s+ال(?:خاص(?:ة|ة)|خصوصية))",
    re.I,
)
HEADING_COMPONENTS = re.compile(
    r"(composantes|areas?\s+of\s+intervention|components?|"
    r"principales?\s+activit[ée]s|main\s+activities|"
    r"م(?:كونات|جالات\s+التدخل)|الأنشطة\s+الرئيسية)",
    re.I,
)
HEADING_KPI = re.compile(
    r"(en\s+chiffres|at\s+a\s+glance|in\s+figures|في\s+أرقام|بالأرقام)",
    re.I,
)
HEADING_KEY_INFO = re.compile(
    r"(informations?\s+cl[ée]s|key\s+information|معلومات\s+أساسية|"
    r"fiche\s+d['']identit[ée]|project\s+factsheet|بطاقة\s+تعريف)",
    re.I,
)
HEADING_HERITAGE = re.compile(
    r"(h[ée]ritage\s+et\s+capitalisation|legacy|إرث\s+و|مكاسب\s+ينبغي)",
    re.I,
)
HEADING_RESULTS = re.compile(
    r"(r[ée]sultats?\s+(?:cl[ée]s|attendus|principaux)|key\s+results|expected\s+results|"
    r"النتائج\s+(?:الرئيسية|المنتظرة))",
    re.I,
)
HEADING_APPROACH = re.compile(
    r"(approche\s+et\s+mise|approach\s+and\s+implementation|المقاربة\s+والتنفيذ|"
    r"une\s+r[ée]ponse\s+aux\s+d[ée]fis|addressing\s+the\s+challenges|الاستجابة\s+لتحديات)",
    re.I,
)
HEADING_PRESENTATION = re.compile(
    r"(pr[ée]sentation\s+du\s+projet|project\s+overview|تقديم\s+المشروع)",
    re.I,
)

PROGRAMME_SECTIONS_FR = [
    (1, "intro"),
    (2, "pourquoi"),
    (3, "vision"),
    (4, "objectifs"),
    (5, "axes"),
    (6, "projets"),
    (7, "territoires"),
    (8, "chiffres"),
    (9, "partenaires"),
    (10, "heritage"),
]

PROJECT_NAMES = [
    "Jeun'ESS",
    "Jeun'ESS",
    "Go4Youth",
    "GO4Youth",
    "SWAFY",
    "Irada4Youth",
    "Maghroum'IN",
    "Fe3il.a",
]


@dataclass
class Para:
    text: str
    style: str


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip())


def paras(doc: Document) -> list[Para]:
    out: list[Para] = []
    for p in doc.paragraphs:
        t = norm(p.text)
        if not t:
            continue
        style = p.style.name if p.style else "normal"
        out.append(Para(t, style))
    return out


def lang_of(text: str) -> str:
    if ARABIC_RE.search(text):
        return "ar"
    if re.search(
        r"\b(the|and|with|for|through|objective|programme|section|young people|"
        r"employment|supporting|project)\b",
        text,
        re.I,
    ):
        return "en"
    return "fr"


def table_kv(table, skip_header: bool = True) -> dict[str, str]:
    rows = table.rows
    start = 1 if skip_header and len(rows) > 1 else 0
    data: dict[str, str] = {}
    for row in rows[start:]:
        cells = [norm(c.text) for c in row.cells]
        if len(cells) >= 2 and cells[0] and cells[1]:
            data[cells[0]] = cells[1]
    return data


def split_lang_blocks(items: list[Para], mode: str) -> dict[str, list[Para]]:
    """Return fr/en/ar paragraph lists."""
    if mode == "fe3ila":
        en_start = next(
            (i for i, p in enumerate(items) if p.text.lower() == "english version"),
            None,
        )
        ar_start = next(
            (i for i, p in enumerate(items) if ARABIC_RE.search(p.text)),
            None,
        )
        fr = items[: en_start or ar_start or len(items)]
        en = items[(en_start + 1 if en_start is not None else 0) : ar_start or len(items)]
        ar = items[ar_start or len(items) :]
        return {"fr": fr, "en": en, "ar": ar}

    if mode == "fr_ar_en":
        first_ar = next((i for i, p in enumerate(items) if ARABIC_RE.search(p.text)), len(items))
        rest = items[first_ar:]
        second_en = None
        for i, p in enumerate(rest):
            if i == 0:
                continue
            if not ARABIC_RE.search(p.text) and lang_of(p.text) == "en":
                if re.search(
                    r"(overall objective|project overview|supporting local|"
                    r"project to promote|project factsheet|gates for)",
                    p.text,
                    re.I,
                ):
                    second_en = first_ar + i
                    break
        if second_en is None:
            for i, p in enumerate(rest):
                if i > 5 and not ARABIC_RE.search(p.text) and re.match(
                    r"^[A-Z][A-Za-z'’\.\-\s]+(?:Youth|ESS|YOUTH|SWAFY|Irada|Fe3il)",
                    p.text,
                ):
                    second_en = first_ar + i
                    break
        fr = items[:first_ar]
        ar = items[first_ar : second_en or len(items)]
        en = items[second_en or len(items) :]
        return {"fr": fr, "en": en, "ar": ar}

    if mode == "programme":
        ar_start = next(
            (i for i, p in enumerate(items) if p.text.startswith("برنامج EU4Youth")),
            None,
        )
        en_start = next(
            (i for i, p in enumerate(items) if "EU4YOUTH TUNISIA PROGRAMME" in p.text.upper()),
            None,
        )
        return {
            "fr": items[: ar_start or en_start or len(items)],
            "ar": items[ar_start or len(items) : en_start or len(items)],
            "en": items[en_start or len(items) :],
        }

    if mode == "eu_tunisie":
        ar_start = next(
            (i for i, p in enumerate(items) if p.text.startswith("الاتحاد الأوروبي")),
            None,
        )
        en_start = next(
            (i for i, p in enumerate(items) if p.text == "The European Union in Tunisia"),
            None,
        )
        return {
            "fr": items[: ar_start or en_start or len(items)],
            "ar": items[ar_start or len(items) : en_start or len(items)],
            "en": items[en_start or len(items) :],
        }

    raise ValueError(f"Unknown split mode: {mode}")


def parse_kpi_line(text: str) -> dict[str, str] | None:
    t = text.rstrip(";").strip()
    m = KPI_RE.match(t)
    if m:
        return {"value": m.group(1).strip(), "label": m.group(2).strip()}
    m2 = re.match(r"^(.+?)\s+(\d[\d\s,\.+]*(?:\s*[%€])?)$", t)
    if m2 and ARABIC_RE.search(t):
        return {"value": m2.group(2).strip(), "label": m2.group(1).strip()}
    if ARABIC_RE.search(t):
        m3 = re.search(r"(\d[\d\s,\.]+)", t)
        if m3:
            value = m3.group(1).strip()
            label = (t[: m3.start()] + t[m3.end() :]).strip(" ;،.")
            if label:
                return {"value": value, "label": label}
    if re.match(r"^[\+\d]", t):
        parts = t.split(None, 1)
        if len(parts) == 2:
            return {"value": parts[0], "label": parts[1]}
    return None


def is_component_heading(text: str) -> bool:
    return bool(
        re.match(
            r"^(\d+[\.\)\-—]\s+|(?:Social|Re-Fund|Market|Community|LIMITL|Jeun|Go4|SWAFY|Irada|Maghroum|Fe3il))",
            text,
            re.I,
        )
        or re.match(r"^\d+[\.\)\-—]\s+", text)
    )


def extract_kpis_from_results(texts: list[str]) -> list[dict[str, str]]:
    kpis: list[dict[str, str]] = []
    for t in texts:
        k = parse_kpi_line(t)
        if k:
            kpis.append(k)
        elif re.search(r"\d+\s", t) and len(t) < 120:
            m = re.match(r"^(\d+)\s+(.+)$", t)
            if m:
                kpis.append({"value": m.group(1), "label": m.group(2)})
    return kpis


def map_key_info(raw: dict[str, str]) -> dict[str, str]:
    mapping = {
        "period": [
            "Période",
            "Durée",
            "Implementation period",
            "Duration",
            "Période d'exécution",
            "فترة التنفيذ",
            "المدة",
        ],
        "implementer": [
            "Organismes de mise en œuvre",
            "Mise en œuvre",
            "Implementing organisations",
            "Implementation",
            "Organisme de mise en œuvre",
            "جهات التنفيذ",
            "الجهة المنفذة",
            "هيئة التنفيذ",
        ],
        "institutionalPartner": [
            "Partenaire institutionnel",
            "Partenaire institutionnel principal",
            "Institutional partner",
            "Main institutional partner",
            "الشريك المؤسسي",
            "الشريك المؤسسي الرئيسي",
        ],
        "targetGroup": [
            "Public cible",
            "Target group",
            "الفئة المستهدفة",
            "Couverture géographique",
            "Geographical coverage",
            "التغطية الجغرافique",
            "التغطية الجغرافية",
        ],
        "funding": [
            "Financement",
            "Funding",
            "التمويل",
        ],
    }
    out = {
        "period": "",
        "implementer": "",
        "institutionalPartner": "",
        "targetGroup": "",
        "funding": "",
    }
    for key, labels in mapping.items():
        for label, val in raw.items():
            if any(l.lower() in label.lower() for l in labels):
                if not out[key]:
                    out[key] = val
    return out


def empty_lang_block() -> dict[str, Any]:
    return {
        "acronym": "",
        "fullName": "",
        "tagline": "",
        "presentation": [],
        "generalObjective": "",
        "specificObjectives": [],
        "kpis": [],
        "components": [],
        "heritage": [],
        "keyInfo": {
            "period": "",
            "implementer": "",
            "institutionalPartner": "",
            "targetGroup": "",
            "funding": "",
        },
    }


def parse_project_lang(block: list[Para]) -> dict[str, Any]:
    data = empty_lang_block()
    if not block:
        return data

    h1s = [p.text for p in block if p.style == "Heading 1"]
    if h1s:
        data["acronym"] = h1s[0]
        if len(h1s) > 1 and len(h1s[1]) > len(h1s[0]) + 5:
            data["fullName"] = h1s[1]

    for p in block[:8]:
        if p.style == "Heading 2" and not data["fullName"]:
            if not HEADING_KPI.search(p.text) and not HEADING_OBJECTIVE.search(p.text):
                if len(p.text) > 40 or "Renforcer" in p.text or "Soutenir" in p.text:
                    if not data["tagline"]:
                        data["tagline"] = p.text
                    elif not data["fullName"]:
                        data["fullName"] = p.text
                elif not data["tagline"]:
                    data["tagline"] = p.text
        elif p.style == "Heading 2" and not data["tagline"] and not HEADING_PRESENTATION.search(p.text):
            if not HEADING_KPI.search(p.text):
                data["tagline"] = p.text
        elif p.style == "Heading 3" and not data["tagline"] and not HEADING_PRESENTATION.search(p.text):
            data["tagline"] = p.text

    section = "intro"
    current_spec: dict[str, str] | None = None
    current_comp: dict[str, Any] | None = None
    comp_mode = False
    result_lines: list[str] = []

    for p in block:
        t = p.text
        if SECTION_SKIP.match(t):
            continue
        if t.lower() == "english version":
            continue
        if re.match(r"^[-—]{5,}$", t):
            continue

        if HEADING_KEY_INFO.search(t):
            section = "keyinfo_marker"
            continue
        if HEADING_KPI.search(t):
            section = "kpis"
            continue
        if HEADING_OBJECTIVE.search(t):
            section = "general"
            continue
        if HEADING_SPECIFIC.search(t):
            section = "specific"
            continue
        if HEADING_RESULTS.search(t):
            section = "results"
            continue
        if HEADING_COMPONENTS.search(t):
            section = "components"
            comp_mode = True
            continue
        if HEADING_APPROACH.search(t):
            section = "tail"
            comp_mode = False
            if current_comp:
                data["components"].append(current_comp)
                current_comp = None
            continue
        if HEADING_HERITAGE.search(t) and not HEADING_PRESENTATION.search(t):
            section = "heritage"
            comp_mode = False
            if current_comp:
                data["components"].append(current_comp)
                current_comp = None
            continue
        if HEADING_PRESENTATION.search(t):
            section = "presentation"
            continue

        if section == "intro" and p.style.startswith("Heading"):
            if p.style == "Heading 1" and t != data["acronym"] and not data["fullName"]:
                data["fullName"] = t
            elif p.style == "Heading 2" and not data["tagline"] and not HEADING_KPI.search(t):
                data["tagline"] = t
            continue

        if section in ("intro", "presentation") and p.style == "normal":
            if not HEADING_KPI.search(t) and not HEADING_OBJECTIVE.search(t):
                data["presentation"].append(t)
            continue

        if section == "kpis":
            k = parse_kpi_line(t)
            if k:
                data["kpis"].append(k)
            continue

        if section == "results" and p.style == "normal":
            result_lines.append(t)
            k = parse_kpi_line(t)
            if k and k not in data["kpis"]:
                data["kpis"].append(k)
            continue

        if section == "general" and p.style == "normal":
            if data["generalObjective"]:
                data["generalObjective"] += " " + t
            else:
                data["generalObjective"] = t
            continue

        if section == "specific":
            if p.style.startswith("Heading"):
                if current_spec and (current_spec.get("body") or current_spec.get("title")):
                    if current_spec["title"] not in ("Résultats attendus", "Expected results", "النتائج المنتظرة"):
                        data["specificObjectives"].append(current_spec)
                title = re.sub(r"^\d+[\.\)]\s*", "", t)
                title = re.sub(r"^Objectif\s+\d+\s*[—\-]\s*", "", title, flags=re.I)
                title = re.sub(r"^Objective\s+\d+\s*[—\-]\s*", "", title, flags=re.I)
                current_spec = {"title": title, "body": ""}
            elif p.style == "normal":
                if current_spec is not None:
                    if current_spec["body"]:
                        current_spec["body"] += " " + t
                    else:
                        current_spec["body"] = t
                elif t.endswith(":") or "objectifs" in t.lower()[:40]:
                    pass
                else:
                    data["specificObjectives"].append(t)
            continue

        if section == "components" or comp_mode:
            if p.style.startswith("Heading") and is_component_heading(t):
                if current_comp:
                    data["components"].append(current_comp)
                current_comp = {"name": t, "description": "", "bullets": []}
                continue
            if p.style.startswith("Heading") and HEADING_RESULTS.search(t):
                section = "results"
                continue
            if current_comp is not None and p.style == "normal":
                if t.lower().startswith(
                    ("objectif :", "objective:", "activités", "main activities", "الهدف:", "الأنشطة")
                ):
                    current_comp["bullets"].append(t)
                elif re.search(r"secteur|sector|القطاع", t, re.I):
                    current_comp["bullets"].append(t)
                elif not current_comp["description"]:
                    current_comp["description"] = t
                else:
                    current_comp["bullets"].append(t)
            continue

        if section == "heritage" and p.style == "normal":
            if not HEADING_OBJECTIVE.search(t):
                data["heritage"].append(t)

    if current_spec and current_spec.get("title") not in ("Résultats attendus", "Expected results"):
        if current_spec.get("body") or current_spec.get("title"):
            data["specificObjectives"].append(current_spec)
    if current_comp:
        data["components"].append(current_comp)

    if not data["kpis"] and result_lines:
        data["kpis"] = extract_kpis_from_results(result_lines)

    data["specificObjectives"] = [
        x
        for x in data["specificObjectives"]
        if not (isinstance(x, dict) and x.get("title") in ("Résultats attendus", "Expected results", "النتائج المنتظرة") and not x.get("body"))
    ]

    if data["tagline"] == data["fullName"] or HEADING_PRESENTATION.search(data.get("tagline") or ""):
        for p in block:
            if p.style == "Heading 2" and not HEADING_KPI.search(p.text) and not HEADING_OBJECTIVE.search(p.text):
                if not HEADING_PRESENTATION.search(p.text) and p.text != data["fullName"]:
                    data["tagline"] = p.text
                    break
        for p in block:
            if p.style == "Heading 3" and not HEADING_PRESENTATION.search(p.text) and not HEADING_KEY_INFO.search(p.text):
                data["tagline"] = p.text
                break

    return data


def assign_tables_project(
    lang_data: dict[str, dict[str, Any]], tables: list, table_order: tuple[str, str, str]
) -> None:
    for ti, lang in enumerate(table_order):
        if ti >= len(tables):
            continue
        raw = table_kv(tables[ti])
        if not raw:
            continue
        ki = map_key_info(raw)
        lang_data[lang]["keyInfo"].update({k: v for k, v in ki.items() if v})
        for label in ("Nom du projet", "Project name", "Acronyme", "Acronym", "اسم المشروع", "الاختصار"):
            if label in raw and not lang_data[lang]["acronym"]:
                lang_data[lang]["acronym"] = raw[label]
        for label in (
            "Nom complet",
            "Full name",
            "الاسم الكامل",
            "Nom court",
            "Short name",
            "الاسم المختصر",
        ):
            if label in raw and not lang_data[lang]["fullName"]:
                lang_data[lang]["fullName"] = raw[label]


def extract_project_doc(path: Path, slug: str, split_mode: str, table_order: tuple[str, str, str]) -> dict:
    doc = Document(path)
    items = paras(doc)
    blocks = split_lang_blocks(items, split_mode)
    result: dict[str, Any] = {"slug": slug}
    for lang in ("fr", "en", "ar"):
        result[lang] = parse_project_lang(blocks.get(lang, []))
    assign_tables_project(result, doc.tables, table_order)
    return result


def join_paragraphs(lines: list[str]) -> str:
    return "\n\n".join(lines)


def parse_programme_lang(block: list[Para]) -> dict[str, dict[str, str]]:
    blocks: dict[str, str] = {}
    if not block:
        return blocks

    current_key = "intro.title"
    buffer: list[str] = []

    def flush():
        nonlocal buffer
        if buffer:
            text = join_paragraphs(buffer)
            if current_key in blocks:
                blocks[current_key] += "\n\n" + text
            else:
                blocks[current_key] = text
            buffer = []

    section_num = 0
    for p in block:
        t = p.text
        if SECTION_SKIP.match(t):
            continue

        sec_m = re.match(r"^SECTION\s+(\d+)", t, re.I)
        if sec_m:
            flush()
            section_num = int(sec_m.group(1))
            continue

        if re.match(r"^القسم\s+(?:ال)?(?:أ|ا)?(?:ول|ثاني|ثالث|رابع|خامس|سادس|سابع|ثامن|تاسع|عاشر)", t):
            flush()
            ar_nums = {
                "الأول": 1,
                "الاول": 1,
                "1": 1,
                "الثاني": 2,
                "الثالث": 3,
                "الرابع": 4,
                "الخامس": 5,
                "السادس": 6,
                "السابع": 7,
                "الثامن": 8,
                "التاسع": 9,
                "العاشر": 10,
            }
            for k, v in ar_nums.items():
                if k in t:
                    section_num = v
                    break
            continue

        if t in ("Surtitre", "Titre principal", "Texte d'introduction"):
            flush()
            key_map = {
                "Surtitre": "intro.eyebrow",
                "Titre principal": "intro.title",
                "Texte d'introduction": "intro.body",
            }
            current_key = key_map[t]
            continue

        if section_num:
            flush()
            name = next((n for num, n in PROGRAMME_SECTIONS_FR if num == section_num), f"section{section_num}")
            if section_num == 1 and not blocks.get("intro.title"):
                if "EU4Youth" in t and len(t) < 200:
                    current_key = "intro.title"
                elif not blocks.get("intro.body"):
                    current_key = "intro.body"
                else:
                    current_key = "intro.body"
            elif section_num == 2:
                if "Pourquoi" in t or "Why" in t or "لماذا" in t:
                    current_key = "pourquoi.title"
                elif t.startswith("«") or t.startswith('"'):
                    current_key = "pourquoi.quote"
                else:
                    current_key = "pourquoi.body"
            elif section_num == 3:
                current_key = "vision.body" if blocks.get("vision.title") else "vision.title"
                if "Vision" in t or "الرؤية" in t:
                    current_key = "vision.title"
            elif section_num == 4:
                if "Objectif général" in t or "Overall objective" in t or "الهدف العام" in t:
                    current_key = "objectifs.general"
                elif re.search(r"objectif\s+sp[ée]cifique|specific objective|الهدف الخصوصي", t, re.I):
                    m = re.search(r"(\d+)", t)
                    n = m.group(1) if m else "1"
                    current_key = f"objectifs.specific{n}"
                else:
                    current_key = "objectifs.general" if "objectif" in t.lower()[:20] else f"{name}.body"
            elif section_num == 5:
                if re.search(r"axe\s+\d|area\s+\d|المحور", t, re.I):
                    m = re.search(r"(\d+)", t)
                    n = m.group(1) if m else "1"
                    current_key = f"axes.axe{n}"
                else:
                    current_key = "axes.intro"
            elif section_num == 6:
                for pn in PROJECT_NAMES:
                    if t.strip().startswith(pn) or t.strip() == pn:
                        slug_key = (
                            pn.lower()
                            .replace("'", "")
                            .replace(".", "")
                            .replace("go4youth", "go4youth")
                        )
                        slug_map = {
                            "jeuness": "jeuness",
                            "go4youth": "go4youth",
                            "swafy": "swafy",
                            "irada4youth": "irada4youth",
                            "maghroumin": "maghroumin",
                            "fe3ila": "fe3ila",
                        }
                        for k, v in slug_map.items():
                            if k in slug_key.replace("'", ""):
                                current_key = f"projets.{v}"
                                buffer = [t.split(None, 1)[1] if " " in t and len(t.split(None, 1)[0]) < 15 else t]
                                continue
                if not current_key.startswith("projets."):
                    if "Six projets" in t or "Six projects" in t or "ستة مشاريع" in t:
                        current_key = "projets.intro"
                    else:
                        current_key = "projets.intro"
            elif section_num == 7:
                current_key = "territoires.body"
            elif section_num == 8:
                if "chiffres" in t.lower() or "figures" in t.lower() or "أرقام" in t:
                    current_key = "chiffres.title"
                elif parse_kpi_line(t) or re.match(r"^[\+\d]", t):
                    k = parse_kpi_line(t) or {"value": t, "label": ""}
                    current_key = f"chiffres.{k['value'].replace('+', 'plus').replace(' ', '_')}"
                else:
                    current_key = "chiffres.body"
            elif section_num == 9:
                if "Union européenne" in t or "European Union" in t or "الاتحاد الأوروبي" in t:
                    current_key = "partenaires.eu"
                elif "institutionnels" in t.lower() or "institutional" in t.lower():
                    current_key = "partenaires.institutionnels"
                elif "mise en œuvre" in t.lower() or "implementing" in t.lower():
                    current_key = "partenaires.implementation"
                elif "acteurs mobilisés" in t.lower() or "local territories" in t.lower():
                    current_key = "partenaires.territoires"
                else:
                    current_key = "partenaires.intro"
            elif section_num == 10:
                current_key = "heritage.body"
            else:
                current_key = f"{name}.body"

        buffer.append(t)

    flush()
    return blocks


def extract_programme(path: Path) -> dict[str, Any]:
    doc = Document(path)
    items = paras(doc)
    blocks = split_lang_blocks(items, "programme")
    return {
        lang: {"blocks": parse_programme_lang(blocks.get(lang, []))}
        for lang in ("fr", "en", "ar")
    }


def split_theme_title_body(text: str) -> tuple[str, str]:
    if ARABIC_RE.search(text):
        # Arabic themes: "Title body..." combined in one paragraph
        themes_ar = [
            "المساواة بين المرأة والرجل",
            "حقوق الإنسان والمجتمع المدني",
            "الصحة",
            "تغير المناخ والطاقة",
            "التنمية الجهوية والمحلية",
            "البيئة والتنمية المستدامة والمياه",
            "الفلاحة",
            "الإعلام والثقافة",
            "التربية والبحث والابتكار",
            "التشغيل والتكوين المهني",
            "الديمقراطية والحوكمة",
            "التنمية الاقتصادية ودعم القطاع الخاص",
        ]
        for title in themes_ar:
            if text.startswith(title):
                return title, text[len(title) :].strip()
        return text.split(None, 2)[0], text
    # English combined lines
    themes_en = [
        "Gender Equality",
        "Human Rights and Civil Society",
        "Health",
        "Climate Change and Energy",
        "Regional and Local Development",
        "Environment, Sustainable Development and Water",
        "Agriculture",
        "Media and Culture",
        "Education, Research and Innovation",
        "Employment and Vocational Training",
        "Democracy and Governance",
        "Economic Development and Private Sector Support",
    ]
    for title in themes_en:
        if text.startswith(title):
            body = text[len(title) :].strip()
            return title, body
    return text, ""


def parse_eu_tunisie_lang(block: list[Para]) -> dict[str, Any]:
    data: dict[str, Any] = {
        "heroTitle": "",
        "intro": "",
        "themesTitle": "",
        "themesLead": "",
        "themes": [],
        "ctaTitle": "",
        "ctaBody": "",
        "links": [],
    }
    if not block:
        return data

    section = "hero"
    current_theme: dict[str, str] | None = None

    for p in block:
        t = p.text
        if p.style.startswith("Heading 1") and not data["heroTitle"]:
            data["heroTitle"] = t
            section = "intro"
            continue
        if section == "intro" and p.style == "normal" and not data["intro"]:
            data["intro"] = t
            continue
        if section == "intro" and p.style == "normal" and data["intro"] and not data.get("_intro2"):
            data["_intro2"] = t
            data["intro"] = join_paragraphs([data["intro"], t])
            section = "themes"
            continue
        if "Thématiques" in t or "Thematic Areas" in t or t == "المحاور":
            data["themesTitle"] = t
            section = "themes"
            continue
        if p.style.startswith("Heading 3"):
            if current_theme:
                data["themes"].append(current_theme)
            current_theme = {"title": t, "body": ""}
            section = "theme_body"
            continue
        if section == "theme_body" and p.style == "normal":
            if current_theme is not None:
                current_theme["body"] = t
                data["themes"].append(current_theme)
                current_theme = None
            continue
        if section == "themes" and p.style == "normal" and not p.style.startswith("Heading"):
            title, body = split_theme_title_body(t)
            if body or len(title) < 80:
                data["themes"].append({"title": title, "body": body or ""})
            continue
        if p.style.startswith("Heading 2") and (
            "Explorer" in t or "Explore European" in t or "مشاريع الاتحاد" in t
        ):
            if current_theme:
                data["themes"].append(current_theme)
                current_theme = None
            data["ctaTitle"] = re.sub(r"\[|\].*", "", t).strip()
            section = "cta"
            continue
        if section == "cta" and p.style == "normal":
            if URL_RE.search(t):
                label = re.sub(r"\[|\]|→.*|https?://.*", "", t).strip()
                for url in URL_RE.findall(t):
                    entry = {"label": label or url, "href": url}
                    if entry not in data["links"]:
                        data["links"].append(entry)
            elif not data["ctaBody"]:
                data["ctaBody"] = t
            else:
                data["ctaBody"] = join_paragraphs([data["ctaBody"], t])

    if current_theme:
        data["themes"].append(current_theme)
    data.pop("_intro2", None)
    return data


def extract_eu_tunisie(path: Path) -> dict[str, Any]:
    doc = Document(path)
    items = paras(doc)
    blocks = split_lang_blocks(items, "eu_tunisie")
    return {lang: parse_eu_tunisie_lang(blocks.get(lang, [])) for lang in ("fr", "en", "ar")}


def write_json(path: Path, payload: Any) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def build_readme(summaries: list[str]) -> str:
    return (
        "# Extracted i18n content\n\n"
        "Structured JSON parsed from trilingual Word source documents in "
        "`EU4Youth/Sections trilingues/` using `tools/extract_trilingual_docx.py` "
        "(python-docx).\n\n"
        "## Files\n\n"
        + "\n".join(f"- {line}" for line in summaries)
        + "\n\n"
        "## Notes\n\n"
        "- Project files (`fe3ila.json`, `go4youth.json`, `irada4youth.json`, `jeuness.json`) "
        "share a common schema with `fr`, `en`, and `ar` blocks.\n"
        "- `programme.json` uses CMS-style semantic keys under `blocks` per language.\n"
        "- `eu-en-tunisie.json` holds hero, themes (~12), CTA, and link metadata.\n"
        "- Key-info tables are mapped as table0=FR, with order varying by document "
        "(see script `table_order` per file).\n"
        "- Encoding: UTF-8.\n"
    )


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    summaries: list[str] = []

    projects = [
        (
            "Fe3ila Trilingue.docx",
            "fe3ila.json",
            "fe3ila",
            "fe3ila",
            ("fr", "en", "ar"),
        ),
        (
            "GO4Youth trilingue.docx",
            "go4youth.json",
            "go4youth",
            "fr_ar_en",
            ("fr", "ar", "en"),
        ),
        (
            "Irada4youth trilingue.docx",
            "irada4youth.json",
            "irada4youth",
            "fr_ar_en",
            ("fr", "ar", "en"),
        ),
        (
            "Jeun_ESS trilingue.docx",
            "jeuness.json",
            "jeuness",
            "fr_ar_en",
            ("fr", "ar", "en"),
        ),
    ]

    for filename, out_name, slug, mode, table_order in projects:
        path = DOC_DIR / filename
        payload = extract_project_doc(path, slug, mode, table_order)
        write_json(OUT_DIR / out_name, payload)
        kpi_counts = {lang: len(payload[lang]["kpis"]) for lang in ("fr", "en", "ar")}
        summaries.append(
            f"`{out_name}` — slug `{slug}`; KPIs fr/en/ar: "
            f"{kpi_counts['fr']}/{kpi_counts['en']}/{kpi_counts['ar']}; "
            f"components fr: {len(payload['fr']['components'])}"
        )

    prog_path = DOC_DIR / "PROGRAMME EU4YOUTH TUNISIE trilingue.docx"
    programme = extract_programme(prog_path)
    write_json(OUT_DIR / "programme.json", programme)
    block_counts = {lang: len(programme[lang]["blocks"]) for lang in ("fr", "en", "ar")}
    summaries.append(
        f"`programme.json` — block keys fr/en/ar: "
        f"{block_counts['fr']}/{block_counts['en']}/{block_counts['ar']}"
    )

    eu_path = DOC_DIR / "Union européenne en Tunisie trilingue.docx"
    eu = extract_eu_tunisie(eu_path)
    write_json(OUT_DIR / "eu-en-tunisie.json", eu)
    theme_counts = {lang: len(eu[lang]["themes"]) for lang in ("fr", "en", "ar")}
    summaries.append(
        f"`eu-en-tunisie.json` — themes fr/en/ar: "
        f"{theme_counts['fr']}/{theme_counts['en']}/{theme_counts['ar']}; "
        f"links fr: {len(eu['fr']['links'])}"
    )

    (OUT_DIR / "README.md").write_text(build_readme(summaries), encoding="utf-8")
    print("Wrote", len(summaries), "JSON files + README to", OUT_DIR)
    for s in summaries:
        print(" ", s)


if __name__ == "__main__":
    main()
