# -*- coding: utf-8 -*-
"""Extract EU4Youth website requirements from source docs."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
import pymupdf

BASE = Path(r"c:\Users\youss\OneDrive\Attachments\Desktop\fi2t-live-editor\EU4Youth\EU4Youth")
SITE = next(p for p in BASE.iterdir() if p.is_dir() and "site web" in p.name.lower())
SITE_INNER = next(p for p in SITE.iterdir() if p.is_dir())
UI = next(p for p in BASE.iterdir() if p.is_dir() and "UI Web Design" in p.name)
UI_INNER = next(p for p in UI.iterdir() if p.is_dir())
OUT = Path(r"c:\Users\youss\OneDrive\Attachments\Desktop\fi2t-live-editor\EU4Youth\_extract_raw.json")


def find_file(root: Path, *parts_or_patterns) -> Path | None:
    """Find file by walking; patterns matched case-insensitively on name."""
    for p in root.rglob("*"):
        if not p.is_file():
            continue
        name = p.name.lower()
        if all(pat.lower() in name for pat in parts_or_patterns):
            return p
    return None


def docx_text(path: Path, max_chars: int | None = None) -> str:
    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    # tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if max_chars and len(text) > max_chars:
        return text[:max_chars] + "\n...[truncated]..."
    return text


def docx_summary(path: Path, max_chars: int = 12000) -> dict:
    text = docx_text(path, max_chars=None)
    # keep headings-ish lines and first portion
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    headings = []
    for ln in lines:
        if len(ln) < 120 and (
            ln.isupper()
            or ln.startswith("#")
            or re.match(r"^(\d+[\.\)]|[A-Z]\.|Section|Page|Chapitre|PARTIE)", ln, re.I)
            or any(
                k in ln.lower()
                for k in [
                    "page",
                    "section",
                    "filtre",
                    "cms",
                    "langue",
                    "mobile",
                    "carte",
                    "map",
                    "accueil",
                    "glossaire",
                    "publication",
                    "opportunité",
                    "actualit",
                    "contact",
                    "projet",
                    "exigence",
                    "fonctionnal",
                    "contenu",
                    "navigation",
                    "footer",
                    "header",
                    "hero",
                    "banner",
                ]
            )
        ):
            headings.append(ln)
    body = "\n".join(lines)
    if len(body) > max_chars:
        body = body[:max_chars] + "\n...[truncated]..."
    return {
        "file": str(path.relative_to(BASE)),
        "chars": len(text),
        "heading_candidates": headings[:80],
        "full_text": body,
    }


def pdf_first_pages(path: Path, n_pages_text: int = 1) -> dict:
    doc = pymupdf.open(str(path))
    page_count = doc.page_count
    texts = []
    for i in range(min(n_pages_text, page_count)):
        texts.append(doc[i].get_text("text"))
    # also try to get section-like titles from all pages (short lines)
    titles = []
    for i in range(page_count):
        t = doc[i].get_text("text")
        for ln in t.splitlines():
            s = ln.strip()
            if not s or len(s) > 80:
                continue
            if s.isupper() or re.match(r"^[A-ZÀÂÄÉÈÊËÎÏÔÙÛÜÇ].{2,}$", s):
                if s not in titles:
                    titles.append(s)
    doc.close()
    first = "\n".join(texts).strip()
    if len(first) > 3000:
        first = first[:3000] + "\n...[truncated]..."
    return {
        "file": str(path.relative_to(BASE)) if path.is_relative_to(BASE) else str(path),
        "page_count": page_count,
        "first_page_text": first,
        "section_titles_visible": titles[:40],
    }


def count_glossary_terms(html_path: Path) -> dict:
    html = html_path.read_text(encoding="utf-8", errors="ignore")
    # common patterns: <dt>, <h2>/<h3>, .term, glossary entries
    patterns = [
        (r"<dt[^>]*>", "dt"),
        (r'class=["\'][^"\']*term[^"\']*["\']', "class_term"),
        (r"<h[23][^>]*>", "h2_h3"),
        (r"<strong[^>]*>[^<]{2,80}</strong>", "strong"),
        (r"<li[^>]*>\s*<[^>]+>[^<]{2,80}</", "li_labeled"),
    ]
    counts = {}
    for pat, name in patterns:
        counts[name] = len(re.findall(pat, html, re.I))
    # try extract term labels from definition lists or bold+def
    terms = re.findall(r"<dt[^>]*>(.*?)</dt>", html, re.I | re.S)
    if not terms:
        terms = re.findall(r"<h3[^>]*>(.*?)</h3>", html, re.I | re.S)
    if not terms:
        terms = re.findall(r"<strong[^>]*>(.*?)</strong>", html, re.I | re.S)
    clean = []
    for t in terms:
        t2 = re.sub(r"<[^>]+>", "", t).strip()
        t2 = re.sub(r"\s+", " ", t2)
        if t2 and len(t2) < 120:
            clean.append(t2)
    # unique preserve order
    seen = set()
    uniq = []
    for t in clean:
        k = t.lower()
        if k not in seen:
            seen.add(k)
            uniq.append(t)
    return {
        "file": str(html_path.relative_to(BASE)),
        "pattern_counts": counts,
        "term_count_estimate": len(uniq),
        "sample_terms": uniq[:30],
        "html_chars": len(html),
    }


def project_one_liners(proj_root: Path) -> list[dict]:
    results = []
    if not proj_root.exists():
        return results
    for folder in sorted([p for p in proj_root.iterdir() if p.is_dir()]):
        docs = list(folder.rglob("*.docx")) + list(folder.rglob("*.pdf"))
        # prefer docx presentation-like
        chosen = None
        for d in docs:
            n = d.name.lower()
            if d.suffix.lower() == ".docx" and any(
                k in n for k in ["présentation", "presentation", "fiche", "projet", "eu4", folder.name.lower()]
            ):
                chosen = d
                break
        if not chosen and docs:
            # first docx else first file
            docxs = [d for d in docs if d.suffix.lower() == ".docx"]
            chosen = docxs[0] if docxs else docs[0]
        entry = {"project_folder": folder.name, "source": None, "summary": None}
        if not chosen:
            entry["summary"] = "(no presentation doc found)"
            results.append(entry)
            continue
        entry["source"] = str(chosen.relative_to(BASE))
        try:
            if chosen.suffix.lower() == ".docx":
                text = docx_text(chosen)
            else:
                pdf = pymupdf.open(str(chosen))
                text = "\n".join(pdf[i].get_text("text") for i in range(min(3, pdf.page_count)))
                pdf.close()
            # compress to one line: name + themes
            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            name = folder.name
            for ln in lines[:15]:
                if folder.name.lower().replace("_", "") in ln.lower().replace(" ", "").replace("_", "") or "eu4" in ln.lower():
                    name = ln[:120]
                    break
            # keywords / themes
            blob = " ".join(lines[:80]).lower()
            theme_keys = [
                "emploi",
                "entrepreneuriat",
                "formation",
                "jeunesse",
                "gouvernance",
                "inclusion",
                "genre",
                "environnement",
                "agriculture",
                "innovation",
                "numérique",
                "digital",
                "mobilité",
                "citoyenneté",
                "ess",
                "économie sociale",
                "climat",
                "éducation",
                "compétences",
                "régions",
                "rural",
                "migration",
                "participation",
                "leadership",
                "startup",
                "pme",
            ]
            themes = [k for k in theme_keys if k in blob]
            # first meaningful paragraph
            para = next((ln for ln in lines if len(ln) > 60 and not ln.isupper()), "")[:220]
            entry["summary"] = f"{name} — themes: {', '.join(themes[:8]) or 'n/a'}; excerpt: {para}"
        except Exception as e:
            entry["summary"] = f"(error reading: {e})"
        results.append(entry)
    return results


def publications_by_project(pub_root: Path) -> dict:
    grouped = defaultdict(list)
    if not pub_root.exists():
        return {}
    for pdf in sorted(pub_root.rglob("*.pdf")):
        # group by first subfolder under pub_root
        try:
            rel = pdf.relative_to(pub_root)
        except ValueError:
            continue
        parts = rel.parts
        group = parts[0] if len(parts) > 1 else "_root"
        grouped[group].append(pdf.name)
    return dict(grouped)


def main():
    # locate priority docs
    files = {
        "cahier_des_charges": find_file(SITE_INNER, "cahier des charges"),
        "cadrage_uxui": find_file(SITE_INNER, "cadrage_uxui"),
        "homepage": find_file(SITE_INNER / "Accueil", "homepage") or find_file(SITE_INNER, "homepage"),
        "a_propos": find_file(SITE_INNER, "propos"),
        "contact_form": find_file(SITE_INNER, "formulaire de contact"),
        "programme_tunisie": find_file(BASE, "programme eu4youth tunisie"),
        "document_sans_titre": find_file(BASE, "document sans titre"),
        "ue_tunisie_docx": find_file(BASE, "union"),
        "ue_tunisie_pdf": None,
        "glossaire_html": find_file(SITE_INNER, "glossaire_eu4youth"),
    }
    for p in BASE.iterdir():
        if p.suffix.lower() == ".pdf" and "union" in p.name.lower():
            files["ue_tunisie_pdf"] = p
        if p.suffix.lower() == ".docx" and "union" in p.name.lower():
            files["ue_tunisie_docx"] = p

    # Prefer matching docx names more carefully
    for p in SITE_INNER.rglob("*.docx"):
        n = p.name.lower()
        if "cahier" in n:
            files["cahier_des_charges"] = p
        elif "cadrage" in n:
            files["cadrage_uxui"] = p
        elif "homepage" in n:
            files["homepage"] = p
        elif "propos" in n:
            files["a_propos"] = p
        elif "formulaire" in n and "contact" in n:
            files["contact_form"] = p

    for p in BASE.glob("*.docx"):
        n = p.name.lower()
        if "programme" in n and "eu4" in n:
            files["programme_tunisie"] = p
        if "sans titre" in n:
            files["document_sans_titre"] = p

    doc_extracts = {}
    for key, path in files.items():
        if path is None:
            doc_extracts[key] = {"error": "not found"}
            continue
        if path.suffix.lower() == ".docx":
            doc_extracts[key] = docx_summary(path, max_chars=14000 if key in ("cahier_des_charges", "cadrage_uxui") else 10000)
        elif path.suffix.lower() == ".pdf":
            doc_extracts[key] = pdf_first_pages(path, n_pages_text=2)
        elif path.suffix.lower() == ".html":
            doc_extracts[key] = count_glossary_terms(path)
        else:
            doc_extracts[key] = {"file": str(path), "note": "unsupported"}

    # UI PDFs
    ui_pdfs = {}
    wanted = [
        "accueil",
        "a propos",
        "actualité",
        "actualite",
        "contact",
        "glossaire",
        "opportunit",
        "carte",
        "publications",
        "projet.pdf",
        "projet banner",
        "color",
        "button",
    ]
    for p in sorted(UI_INNER.glob("*.pdf")):
        n = p.name.lower()
        # also root copies
        label = p.stem
        # extract more pages for map/guidelines
        n_text = 2 if any(k in n for k in ["carte", "color", "button", "projet"]) else 1
        ui_pdfs[label] = pdf_first_pages(p, n_pages_text=n_text)

    # root projet pdfs
    for p in BASE.glob("*.pdf"):
        if "projet" in p.name.lower():
            ui_pdfs[p.stem + " (root)"] = pdf_first_pages(p, n_pages_text=2)

    # projects
    proj_root = SITE_INNER / "Projets"
    if not proj_root.exists():
        for p in SITE_INNER.iterdir():
            if p.is_dir() and p.name.lower().startswith("projet"):
                proj_root = p
    projects = project_one_liners(proj_root)

    # publications
    pub_root = None
    for p in SITE_INNER.iterdir():
        if p.is_dir() and "publication" in p.name.lower():
            pub_root = p
            break
    pubs = publications_by_project(pub_root) if pub_root else {}

    # also list all project docx names for transparency
    project_files = []
    if proj_root.exists():
        for p in sorted(proj_root.rglob("*")):
            if p.is_file() and p.suffix.lower() in {".docx", ".pdf", ".pptx"}:
                project_files.append(str(p.relative_to(BASE)))

    result = {
        "paths": {
            "BASE": str(BASE),
            "SITE_INNER": str(SITE_INNER),
            "UI_INNER": str(UI_INNER),
            "resolved_files": {k: (str(v) if v else None) for k, v in files.items()},
            "pub_root": str(pub_root) if pub_root else None,
            "proj_root": str(proj_root),
        },
        "documents": doc_extracts,
        "design_comps_raw": ui_pdfs,
        "projects": projects,
        "project_files": project_files,
        "publications_by_project": pubs,
        "glossaire": doc_extracts.get("glossaire_html"),
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUT}")
    print("Resolved:")
    for k, v in files.items():
        print(f"  {k}: {v.name if v else None}")
    print(f"UI PDFs: {len(ui_pdfs)}")
    print(f"Projects: {len(projects)}")
    print(f"Pub groups: {list(pubs.keys())}")


if __name__ == "__main__":
    main()
